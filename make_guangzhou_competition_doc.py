from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "Guangzhou_Innovation_Competition_TLaser_2026_MaRong.docx"


def set_run_font(run, size=11, bold=False, color=None, font="Microsoft YaHei"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_para(doc, text="", size=11, bold=False, color=None, align=None, after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    for i, part in enumerate(text.split("\n")):
        if i:
            p.add_run().add_break()
        run = p.add_run(part)
        set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_heading(doc, text, level=1):
    size = {1: 16, 2: 13, 3: 12}[level]
    color = "2E74B5" if level in (1, 2) else "1F4D78"
    p = add_para(doc, text, size=size, bold=True, color=color, after=5)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    return p


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=10.5, bold=bold)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "left", "bottom", "right"):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), "100")
        node.set(qn("w:type"), "dxa")


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            set_cell_text(table.cell(r, c), text, bold=(r == 0 or c == 0))
    return table


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.8)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

styles = doc.styles
styles["Normal"].font.name = "Microsoft YaHei"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(11)

add_para(
    doc,
    "2026年广州科技创新创业大赛行业赛\n项目报名材料草案",
    size=22,
    bold=True,
    color="0B2545",
    align=WD_ALIGN_PARAGRAPH.CENTER,
    after=4,
)
add_para(
    doc,
    "项目名称：光通讯激光器数字孪生技术（TLaser）",
    size=13,
    bold=True,
    align=WD_ALIGN_PARAGRAPH.CENTER,
)
add_para(
    doc,
    "拟申报通道：广州市科学技术局主办、 大湾区科技创新服务中心承办的2026年广州科技创新创业大赛行业赛；粤港澳大湾区国家技术创新中心作为特别支持单位和产业协同平台。",
    size=10.5,
    align=WD_ALIGN_PARAGRAPH.CENTER,
)

add_heading(doc, "一、申报通道判断", 1)
add_table(
    doc,
    [
        ["项目", "建议填写/判断"],
        ["受理/主办机构", "广州市科学技术局"],
        ["执行/服务机构", "大湾区科技创新服务中心，广州科技创新创业大赛组委会"],
        ["特别支持平台", "粤港澳大湾区国家技术创新中心"],
        ["建议参赛领域", "战略性新兴产业领域：软件与互联网 / 新一代信息设备和通信网络 / 半导体与集成电路 / 人工智能"],
        ["优先赛道建议", "首选“软件与互联网”或“新一代信息设备和通信网络”；备用“半导体与集成电路”"],
        ["主体提醒", "战略性新兴产业领域通常面向广州市注册企业；如当前主体尚未在广州注册，需尽快确认是否以拟落地广州/拟注册主体报名或转走其他开放赛道。"],
    ],
)

add_heading(doc, "二、项目摘要", 1)
add_para(
    doc,
    "TLaser 是面向光通讯半导体激光器的实时数字孪生与在线标定工业软件。项目聚焦光模块、激光芯片和测试装备企业在研发与测试中的关键痛点：离线仿真耗时、L-I-V 测试数据不能及时反哺设计模型、实测曲线偏离原因难解释、器件性能筛选依赖专家经验。TLaser 以腔长、脊/有源区宽度、有源区厚度、注入电流、热沉/环境温度、前镜反射率、后镜反射率等 7 个设计/工况参数为输入，毫秒级预测 L-I-V 曲线、输出光功率、阈值电流、斜率效率和转换效率；同时接入客户给定的 L-I-V 实测数据，反演内部物性参数并生成校准后的设备级数字孪生。项目当前已完成准三维激光器物理模拟器、PINN 代理模型、参数标定脚本、中英文控制面板、用户手册和演示材料，可进入真实客户数据联合验证和广州/大湾区产业落地阶段。",
)

add_heading(doc, "三、核心技术与创新性", 1)
for item in [
    "7D 参数化激光器物理模型：把工程师可直接理解的设计/工况参数映射到 L-I-V 曲线和关键性能指标。",
    "PINN 代理模型：把物理约束嵌入神经网络训练，使网页端参数扫描达到毫秒级响应。",
    "L-I-V 在线标定算法：根据客户上传的实测曲线反演内损耗、限制因子、串联/分流电阻、非辐射复合等漂移或未知参数。",
    "SaaS 工程工作流：把设计探索、数据上传、自动标定、误差评价、模型保存和报告导出封装为轻量级工业软件服务。",
    "竞争定位：不替代 Ansys、RSoft、Crosslight 等完整 TCAD/光学仿真工具，而是成为重型仿真软件与 L-I-V 测试设备之间的在线标定和即时预测层。",
]:
    add_para(doc, f"• {item}", after=3)

add_heading(doc, "四、产品形态", 1)
add_table(
    doc,
    [
        ["模块", "功能", "客户价值"],
        ["7D 设计探索", "输入或滑动调整 7 个设计/工况参数，实时生成性能预测", "快速筛选设计参数，减少离线仿真和样品试错"],
        ["实测数据标定", "上传 CSV/JSON 格式 L-I-V 曲线，自动清洗、拟合和反演参数", "解释实测偏差，形成校准后的设备级模型"],
        ["器件模型库", "保存器件型号、测试条件、标定参数和版本记录", "让研发/测试团队复用历史模型和客户数据资产"],
        ["报告与 API", "输出校准前后对比、误差指标和工程验证报告", "便于接入测试台、研发数据库和 PLM 流程"],
    ],
)

add_heading(doc, "五、团队与分工", 1)
add_table(
    doc,
    [
        ["成员", "角色", "核心贡献"],
        ["万振文", "CEO/CTO，项目负责人", "中科院博士、国家级人才，长期从事计算物理、CAE/工业软件算法和客户交付；负责核心算法、产品定义、研发组织和重点客户交付。"],
        ["马戎", "商务负责人/COO，拟全职", "深圳市光彩凯宜电子开发有限公司法定代表人、执行董事&总经理；长期服务华为、H3C、浪潮信息、曙光信息等 ICT 客户，负责深圳/华南/大湾区客户开发、试点合作、销售渠道、合同谈判和测试装备伙伴导入。"],
        ["蒋寻涯", "科学顾问，兼职", "复旦大学教授/博导，光电仿真和 EastWave 产业化经验；负责光电仿真技术路线、L-I-V 测试验证方法和种子客户资源协同。"],
    ],
)

add_heading(doc, "六、市场定位与广州/大湾区落地", 1)
add_para(
    doc,
    "TLaser 的目标客户是光模块企业、激光芯片/泵浦模块企业、测试装备企业、硅光/光电芯片企业和研发测试服务机构。广州与大湾区具有通信设备、ICT 终端、智能制造、半导体和软件服务产业基础，深圳/东莞/广州/佛山等城市也有大量测试装备、电子制造和光通信产业链资源。项目在广州参赛时应突出“AI+工业软件”“新一代信息设备和通信网络”“半导体器件研发工具”三重属性，并把马戎的深圳/华南客户资源作为大湾区市场进入抓手。",
)

add_heading(doc, "七、商业模式与三年目标", 1)
add_table(
    doc,
    [
        ["阶段", "产品/客户目标", "收入模式"],
        ["第一年", "完成 V1.0 工程版，接入至少 1 家客户真实 L-I-V 数据并完成校准验证", "联合验证与模型适配服务，单项目收费"],
        ["第二年", "形成 SaaS 试点版、器件模型库和私有化部署包，完成 3 家以上付费试点", "SaaS 订阅 + 模型适配 + 私有化部署试点"],
        ["第三年", "形成商业版、API 接口和标准交付流程，服务不少于 6 类电信激光器型号", "年度授权、私有化部署、API 授权和持续模型服务"],
    ],
)
add_para(
    doc,
    "三年目标：新增项目经费约 800 万元，主营业务收入约 1330 万元，新增就业 25 人，受理发明专利 5 项、PCT 2 项、软件著作权 6 项。广州/大湾区版本可将资金用途表述为：产品工程化、真实器件测试数据接入、云端平台、客户联合验证、知识产权布局和深圳/广州市场拓展。",
)

add_heading(doc, "八、报名材料清单", 1)
for item in [
    "广州科技 GI 账号注册与主体认证信息。",
    "企业营业执照或拟落地主体说明；若战略性新兴产业赛道要求广州注册企业，需尽快确认报名主体。",
    "项目商业计划书或报名表摘要，可直接拆用本文内容。",
    "核心团队简历：万振文、马戎、蒋寻涯。",
    "知识产权与技术基础材料：已授权专利、软件原型、用户手册、演示视频、测试/验证说明。",
    "路演材料：建议同步制作 10 页以内 PPT，突出痛点、产品、技术、客户、商业模式和广州/大湾区落地计划。",
]:
    add_para(doc, f"• {item}", after=3)

doc.save(OUT)
print(f"saved={OUT}")
