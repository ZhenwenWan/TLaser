from pathlib import Path

from docx import Document


SRC = Path("SuzhouIndustrialPark_TLaser_revised_0813.docx")
TEMPLATE = Path("ModuleNew.docx")
OUT = Path("ModuleNew_TLaser_filled.docx")


def cell_text(table, row, col=0):
    return table.rows[row].cells[col].text


def set_cell(cell, text):
    cell.text = text


def replace_paragraph_exact_or_after_heading(doc, heading, text):
    for i, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip() == heading:
            if i + 1 < len(doc.paragraphs):
                doc.paragraphs[i + 1].text = text
                return True
    raise ValueError(f"Heading not found: {heading}")


def copy_table_text(src_table, dst_table):
    rows = min(len(src_table.rows), len(dst_table.rows))
    cols = min(len(src_table.columns), len(dst_table.columns))
    for r in range(rows):
        for c in range(cols):
            set_cell(dst_table.rows[r].cells[c], src_table.rows[r].cells[c].text)


src = Document(SRC)
dst = Document(TEMPLATE)

# Most forms share the same table geometry. Copy filled cells where the template
# did not change its internal row structure.
for ti in [0, 1, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18, 19]:
    copy_table_text(src.tables[ti], dst.tables[ti])

# ModuleNew splits the project summary into seven labeled fields. Reuse the
# audited TLaser SaaS-scope content and route it into the new labels.
summary = dst.tables[3]

project_basis = (
    "项目前期工作基础：TLaser 已完成光通讯半导体激光器准三维物理模拟、PINN 代理模型、"
    "L-I-V 参数标定脚本、中英文 Streamlit 控制面板、自动化验证流程、用户手册和演示材料，"
    "具备接入真实 L-I-V 数据开展客户联合验证的基础。\n"
    "主攻方向：聚焦光通讯半导体激光器研发与测试环节，建设“7D设计空间即时性能预测 + "
    "实测 L-I-V 数据在线标定”的 SaaS 工业软件平台。\n"
    "目标：以腔长、脊/有源区宽度、有源区厚度、注入电流、热沉/环境温度、前镜反射率、"
    "后镜反射率等 7 个设计/工况选项为核心输入，毫秒级预测 L-I-V 曲线、输出光功率、"
    "阈值电流、斜率效率和转换效率；同时接入客户给定 L-I-V 测试数据，反演内部物性参数并"
    "生成校准后的设备级数字孪生。\n"
    "拟解决问题：解决传统激光器研发中离线仿真耗时、测试数据不能及时反哺模型、实测曲线"
    "偏离原因难解释、设计筛选和测试标定脱节等问题。"
)

team_advantage = cell_text(src.tables[3], 1)

core_tech = (
    "项目核心技术是否为卡脖子技术：☑是 □否\n"
    "项目核心技术是否为颠覆性技术：☑是 □否\n"
    + cell_text(src.tables[3], 3)
)

target_product = cell_text(src.tables[3], 5)
maturity = cell_text(src.tables[3], 8)
marketing = cell_text(src.tables[3], 10)
three_years = cell_text(src.tables[3], 12)

summary.rows[1].cells[0].text = project_basis
summary.rows[3].cells[0].text = team_advantage
summary.rows[5].cells[0].text = "项目核心技术是否为卡脖子技术：☑是 □否\n项目核心技术是否为颠覆性技术：☑是 □否\n（如选择“是”，请详细描述）"
summary.rows[6].cells[0].text = core_tech
summary.rows[8].cells[0].text = "产品形态：□最终消费产品 ☑工业产品 □工业中间产品 ☑技术服务 ☑其他（SaaS/私有化工业软件）"
summary.rows[9].cells[0].text = target_product
summary.rows[11].cells[0].text = "项目所属阶段：□论证 ☑研发 □中试 □批量生产"
summary.rows[12].cells[0].text = maturity
summary.rows[14].cells[0].text = marketing
summary.rows[16].cells[0].text = three_years

# Main plan paragraphs in ModuleNew have guidance text immediately after each
# heading. Replace those guidance paragraphs with the revised TLaser content.
src_plan = {
    "1、项目概述": src.paragraphs[63].text,
    "2、公司概况": src.paragraphs[65].text,
    "1、技术研究内容": src.paragraphs[68].text,
    "2、产品与产业化基础": src.paragraphs[70].text,
    "1、市场定位": src.paragraphs[74].text,
    "2、营销策略": src.paragraphs[76].text,
    "1、阶段目标(3～5年)": src.paragraphs[79].text,
    "2、长期目标(6～10年)": src.paragraphs[81].text,
    "1、财务情况": src.paragraphs[84].text,
    "2、融资情况": src.paragraphs[86].text,
    "3、风险预估": src.paragraphs[88].text,
}
for heading, text in src_plan.items():
    replace_paragraph_exact_or_after_heading(dst, heading, text)

dst.save(OUT)
print(f"Saved {OUT}")
