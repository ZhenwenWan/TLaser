from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.shared import Pt


PATH = Path(r"C:\Users\aw4wz\Documents\Codex\TLaser\ModuleNew_0814_LiYingsong_smallteam.docx")


def set_font(run, size=10.5):
    run.font.name = "仿宋"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "仿宋")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)


def replace_cell(cell, text):
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    p = cell.paragraphs[0]
    p.clear()
    run = p.add_run(text)
    set_font(run)


def set_row_values(row, mapping):
    for idx, value in mapping.items():
        if idx < len(row.cells):
            replace_cell(row.cells[idx], value)


doc = docx.Document(PATH)
t15 = doc.tables[15]
set_row_values(t15.rows[13], {0: "提交前补录", 1: "提交前补录"})
set_row_values(t15.rows[14], {0: "提交前补录", 1: "提交前补录", 2: "提交前补录", 3: "提交前补录"})
doc.save(PATH)
print(PATH)
