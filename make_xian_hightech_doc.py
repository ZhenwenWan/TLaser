from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "Xian_HighTech_TLaser_2026_QinSantuan.docx"


def set_run_font(run, size=11, bold=False, color=None, font="Microsoft YaHei"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_para_format(p, after=8, before=0, line=1.25, align=None):
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = line
    if align is not None:
        p.alignment = align


def add_para(doc, text="", size=11, bold=False, color=None, align=None, after=8, before=0):
    p = doc.add_paragraph()
    set_para_format(p, after=after, before=before, align=align)
    for i, part in enumerate(text.split("\n")):
        if i:
            p.add_run().add_break()
        run = p.add_run(part)
        set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_heading(doc, text, level=1):
    size = {1: 16, 2: 13, 3: 12}[level]
    color = "2E74B5" if level in (1, 2) else "1F4D78"
    p = add_para(doc, text, size=size, bold=True, color=color, before=16 if level == 1 else 10, after=6)
    return p


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, margin=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "left", "bottom", "right"):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin))
        node.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, bold=False, fill=None):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    if fill:
        shade_cell(cell, fill)
    p = cell.paragraphs[0]
    set_para_format(p, after=0, line=1.15)
    run = p.add_run(text)
    set_run_font(run, size=10.5, bold=bold)


def set_table_geometry(table, widths):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_grid = tbl.tblGrid
    for grid_col in list(tbl_grid):
        tbl_grid.remove(grid_col)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")


def add_table(doc, rows, widths):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            set_cell_text(table.cell(r, c), text, bold=(r == 0 or c == 0), fill="F4F6F9" if r == 0 else None)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style=None)
        set_para_format(p, after=4, line=1.208)
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.194)
        run = p.add_run("• ")
        set_run_font(run, size=11)
        run = p.add_run(item)
        set_run_font(run, size=11)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1.0)
section.bottom_margin = Inches(1.0)
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
styles["Normal"].font.name = "Microsoft YaHei"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(11)

add_para(
    doc,
    "西安高新区/陕西赛区预申报材料草案\n光通讯激光器数字孪生技术（TLaser）",
    size=22,
    bold=True,
    color="0B2545",
    align=WD_ALIGN_PARAGRAPH.CENTER,
    after=4,
)
add_para(
    doc,
    "申报定位：面向西安高新区政策申报平台、科技型中小企业评价、陕西赛区创新创业大赛及后续硬科技项目征集的备投版材料",
    size=11,
    bold=True,
    align=WD_ALIGN_PARAGRAPH.CENTER,
    after=8,
)
add_para(
    doc,
    "版本日期：2026年8月13日；项目名称采用此前版本名称“光通讯激光器数字孪生技术”。",
    size=10.5,
    color="555555",
    align=WD_ALIGN_PARAGRAPH.CENTER,
)

add_heading(doc, "一、当前申报口判断", 1)
add_para(
    doc,
    "截至2026年8月13日公开检索，西安高新区信用金融服务平台仍是政策申报“全流程”线上办理服务窗口；但可见的2026年第二批高层次人才认定窗口为2026年7月6日至7月19日18:00，已在今日之前截止。暂未发现西安高新区仍在开放的、明确面向TLaser这类企业项目的科技创新专项资金直接申报口。建议把本材料作为西安高新区政策平台建档、科技型中小企业评价、陕西赛区创新创业大赛及后续硬科技项目征集的预申报文本。",
)
add_table(
    doc,
    [
        ["入口/事项", "当前状态判断", "TLaser建议动作"],
        ["西安高新区信用金融服务平台（政策申报—政策宝）", "官方线上申报入口持续存在，可用于政策匹配、申报进度和后续项目通知跟踪。", "先完成企业信息、团队信息、知识产权和产品材料归档，等待新一批科技创新专项、硬科技人才或产业扶持入口。"],
        ["2026年第二批西安高新区高层次人才认定", "公开通知显示申报期为2026年7月6日—7月19日18:00，当前已截止。", "不作为本次立即提交口；可按其材料要求提前准备劳动/服务协议、学历学位、业绩佐证和人才材料。"],
        ["科技型中小企业评价", "公开政策服务信息显示2026年度评价仍可准备，通常是后续政策、比赛和资质申报的基础条件。", "若已有陕西/西安主体，优先完成科技型中小企业评价；若尚未设立主体，先确定拟落地主体和注册地。"],
        ["第十五届中国创新创业大赛陕西赛区", "西安高新区科技创新局已出现陕西赛区会务服务采购信息，说明赛事组织正在筹备；未检索到参赛企业正式报名通知。", "按陕西赛区历史规则预先准备报名表、商业计划书、路演PPT和财务预测，等待官方报名入口开放。"],
    ],
    [2100, 3300, 3960],
)

add_heading(doc, "二、项目摘要", 1)
add_para(
    doc,
    "TLaser是面向光通讯半导体激光器的实时数字孪生与在线标定SaaS平台。项目聚焦电信级边发射激光器、泵浦激光器和相关光模块研发测试中的高频痛点：传统TCAD/光电热仿真成本高、等待时间长；L-I-V实测曲线难以及时反馈到设计模型；工程师依赖经验筛选器件和解释偏差；客户测试数据无法沉淀为可复用的器件级模型。",
)
add_para(
    doc,
    "TLaser以7D设计/工况参数为输入，覆盖腔长、脊/有源区宽度、有源区厚度、注入电流、热沉或环境温度、前镜反射率和后镜反射率，毫秒级预测L-I-V曲线、输出光功率、阈值电流、斜率效率和转换效率；同时支持客户上传给定L-I-V实测数据，自动反演内部物性参数并生成校准后的器件数字孪生模型。项目已形成准三维激光器物理模拟器、PINN代理模型、参数标定脚本、中英文控制面板、用户手册和演示材料，可进入真实客户数据联合验证与工程化交付阶段。",
)

add_heading(doc, "三、技术创新与西安适配性", 1)
add_bullets(
    doc,
    [
        "AI+工业软件：用PINN代理模型把重型物理仿真的响应时间压缩到网页端交互级，支撑研发工程师快速扫描设计空间。",
        "器件级数字孪生：不是做通用光学仿真软件，而是聚焦电信二极管激光器的性能预测、数据标定和模型复用。",
        "测试数据闭环：把客户L-I-V曲线接入模型反演流程，形成可解释的参数漂移诊断和校准后预测。",
        "西安产业适配：西安高新区具有硬科技、集成电路、电子信息、通信测试和高校科研资源基础，TLaser可作为光电子器件研发、通信设备测试和工业软件国产化方向的交叉项目。",
        "与秦三团团队协同：秦三团的射频电路、高速数字电路、FPGA/ARM嵌入式软件和通信系统研发背景，可补强激光器驱动、L-I-V测试硬件、数据采集接口和客户现场验证。",
    ],
)

add_heading(doc, "四、产品形态", 1)
add_table(
    doc,
    [
        ["模块", "主要功能", "客户价值"],
        ["7D设计探索", "输入或滑动调整7个设计/工况参数，实时生成性能预测。", "快速筛选设计参数，减少离线仿真和样品试错。"],
        ["实测数据标定", "上传CSV/JSON格式L-I-V曲线，自动清洗、拟合和反演参数。", "解释实测偏差，形成校准后的器件级模型。"],
        ["器件模型库", "保存器件型号、测试条件、标定参数和版本记录。", "让研发/测试团队复用历史模型和客户数据资产。"],
        ["报告与API", "输出校准前后对比、误差指标和工程验证报告。", "便于接入测试台、研发数据库、PLM或客户私有系统。"],
    ],
    [1900, 3800, 3660],
)

add_heading(doc, "五、团队与西安资源", 1)
add_table(
    doc,
    [
        ["成员", "角色", "核心贡献"],
        ["万振文", "项目负责人/CEO/CTO", "中科院博士，长期从事计算物理、CAE/工业软件算法和客户交付；负责核心算法、产品定义、研发组织、重点客户验证和融资申报。"],
        ["秦三团", "西安技术与测试验证负责人", "西安电子科技大学电子科学与技术博士，现为西安邮电大学导师，曾任中兴通讯3G/4G/5G移动通信系统研发工程师、整机系统研发工程师和项目技术总监；负责激光器驱动与L-I-V测试硬件、数据采集接口、FPGA/ARM嵌入式协同、测试平台联调和西安高校/产业资源对接。"],
        ["马戎", "商务负责人/COO，拟全职", "具备ICT客户服务和商务拓展经验，负责光模块、通信设备、测试设备客户开发，推动西安与华南客户试点、合同谈判和渠道合作。"],
    ],
    [1500, 2500, 5360],
)

add_heading(doc, "六、拟申报方向与落地方案", 1)
add_bullets(
    doc,
    [
        "主申方向：硬科技创新、工业软件、人工智能赋能制造、新一代信息技术和光电子器件研发工具。",
        "主体策略：若已有陕西/西安注册主体，优先用该主体完成政策平台信息维护和科技型中小企业评价；若尚未设立，可将本材料作为拟落地西安高新区的项目对接材料。",
        "合作载体：依托西安高校电子信息、通信测试和嵌入式系统资源，建立L-I-V测试数据采集、激光器驱动板卡和软件标定平台的联合验证。",
        "示范客户：优先选择光模块、激光器芯片、测试设备、通信系统研发单位，围绕1-2类典型电信激光器完成真实数据校准。",
        "提交节奏：短期投科技型中小企业评价和政策平台建档；中期投陕西赛区创新创业大赛；后续根据西安高新区科技创新专项、硬科技人才或秦创原相关通知补齐正式申报表。",
    ],
)

add_heading(doc, "七、商业模式与三年目标", 1)
add_table(
    doc,
    [
        ["阶段", "产品/客户目标", "收入模式"],
        ["第一年", "完成V1.0工程版，接入至少1家真实客户L-I-V数据并完成标定验证。", "联合验证、模型适配服务、原型部署。"],
        ["第二年", "形成SaaS试点版、器件模型库和私有化部署包，完成3家以上付费试点。", "SaaS订阅、模型适配、私有化部署试点。"],
        ["第三年", "形成商业版、API接口和标准交付流程，服务不少于6类电信激光器型号。", "年度授权、API授权、持续模型服务和私有化部署。"],
    ],
    [1500, 4560, 3300],
)
add_para(
    doc,
    "三年目标：新增项目经费约800万元，主营业务收入约1330万元，新增就业25人，受理发明专利5项、PCT 2项、软件著作权6项。西安版资金用途建议表述为：产品工程化、真实器件测试数据接入、云端平台与私有化部署、激光器测试硬件协同、客户联合验证、知识产权布局和西安高新区市场拓展。",
)

add_heading(doc, "八、拟提交材料清单", 1)
add_bullets(
    doc,
    [
        "企业主体材料：营业执照、税务登记/纳税记录、社保或人员证明、公司章程、股权结构。",
        "项目材料：项目摘要、商业计划书、技术路线图、产品演示链接、用户手册、软件原型截图、测试验证说明。",
        "团队材料：万振文、秦三团、马戎简历；秦三团学历、工作经历、科研项目和测试平台资源佐证。",
        "知识产权材料：已授权专利、拟申请专利清单、软件著作权规划、核心算法和数据安全说明。",
        "申报策略材料：科技型中小企业评价资料、陕西赛区创新创业大赛路演PPT、西安高新区政策平台项目介绍、拟落地或合作协议草案。",
    ],
)

add_heading(doc, "九、申报风险提示", 1)
add_para(
    doc,
    "本版材料不假定西安高新区当前存在仍开放的科技专项资金直投入口。若正式申报通知要求企业已在西安高新区注册、纳税或有固定研发场地，应先完成主体、场地、人员和社保/个税安排，再提交资金类项目。若走陕西赛区创新创业大赛，应等官方参赛通知发布后对照行业分组、参赛资格、成立年限、融资轮次和财务要求更新表格。秦三团信息已按公开高校导师页面整理，正式提交前需由本人确认职务、参与方式、工作量和授权使用个人履历材料。",
)

add_heading(doc, "十、参考公开来源", 1)
add_bullets(
    doc,
    [
        "西安高新区信用金融服务平台：政策申报“全流程”线上办理服务窗口。",
        "《关于开展2026年第二批西安高新区高层次人才认定工作的通知》：申报时间为2026年7月6日—7月19日18:00，申报网址为西安高新区信用金融服务平台（政策申报—政策宝）。",
        "中国政府采购网《第十五届中国创新创业大赛陕西赛区会务服务采购项目》：采购人为西安高新技术产业开发区科技创新局，用于判断陕西赛区赛事组织正在筹备。",
        "西安邮电大学研究生院秦三团导师页面：用于整理秦三团教育背景、通信系统研发经历、实验室和研究方向。",
    ],
)

doc.save(OUT)
print(f"saved={OUT}")
