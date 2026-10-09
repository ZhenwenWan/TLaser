from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


SRC = Path(r"C:\Users\aw4wz\Documents\Codex\TLaser\SuzhouIndustrialPark_TLaser_revised_0811.docx")
OUT = Path(r"C:\Users\aw4wz\Documents\Codex\TLaser\SuzhouIndustrialPark_TLaser_revised_0811_JiangAdvisor.docx")


def set_run_font(run, size=10.5):
    run.font.name = "SimSun"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)


def set_paragraph_text(p, text):
    for r in list(p.runs):
        p._p.remove(r._r)
    for i, part in enumerate(text.split("\n")):
        if i:
            p.add_run().add_break()
        run = p.add_run(part)
        set_run_font(run)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(3)


def set_cell_text(cell, text):
    cell.text = ""
    set_paragraph_text(cell.paragraphs[0], text)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "left", "bottom", "right"):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), "100")
        node.set(qn("w:type"), "dxa")


def replace_para_start(doc, prefix, text):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            set_paragraph_text(p, text)
            return True
    return False


def table_text(table):
    return "\n".join(cell.text for row in table.rows for cell in row.cells)


def delete_row(table, row_idx):
    row = table.rows[row_idx]
    tr = row._tr
    tr.getparent().remove(tr)


doc = Document(str(SRC))

summary = (
    "1）团队情况：项目由万振文全职牵头。万振文为中科院博士、国家级人才，曾任厦门大学副教授、丹麦气象研究所资深研究员，近十年从事 CAE/工业软件算法开发，完成 EML、MySim 等软件原型和多项客户合同交付，已取得 3 项发明专利授权。蒋寻涯教授为复旦大学研究员、博士生导师，曾入选中科院百人计划，拟兼职加盟，任科学顾问，负责光电仿真技术路线、L-I-V 测试验证路线和种子客户协调对接。项目早期采用小团队模式，先由万振文负责产品和算法，由蒋寻涯负责顾问支持和客户资源协调，再按试点进展招聘软件、应用工程和销售人员。"
    "\n2）目标产品：TLaser 是面向光通讯半导体激光器的数字孪生与在线标定平台，包含激光器参数快速仿真、L-I-V 实测曲线反演、批次漂移诊断、良率分析和云端模型管理模块。"
    "\n3）核心技术：项目以准三维激光器物理模型生成训练数据，用 PINN 代理模型进行快速预测，再用实测 L-I-V 数据反向校准内部物性参数。模型输入包括反射率、腔长、温度、电流、有源区宽度和厚度等参数，输出包括全局性能指标以及沿腔长方向的载流子密度和光功率剖面估计。"
    "\n4）应用场景：平台面向光模块、激光芯片和测试装备企业，用于器件设计筛选、失效定位、量产一致性分析和测试数据建模。工程验证主要通过 L-I-V 曲线、输出光功率、光谱、近/远场、温度扫描和调制响应等可测量数据进行校准与一致性检验。"
    "\n5）产业化基础：项目已完成模拟器、1500 组训练样本、PINN 模型、参数标定脚本和中英文控制面板，可进入真实客户数据接入和联合验证阶段。目标客户优先覆盖苏州园区及苏州光通信产业链企业，并逐步拓展至长三角客户。"
    "\n6）商业模式与三年目标：第一阶段按联合验证和模型适配服务收费，第二阶段形成 SaaS 订阅，第三阶段提供私有化部署和 API 授权。三年计划新增项目经费 800 万元，实现主营收入 1330 万元，新增就业 25 人，受理发明专利 5 项，登记软件著作权 6 项。"
)
set_paragraph_text(doc.paragraphs[18], summary)

t3 = doc.tables[3]
set_cell_text(
    t3.cell(3, 0),
    "万振文：拟全职，CEO/CTO，负责公司战略、核心算法、产品定义、研发组织和重点客户交付。其工作经历覆盖科研建模、工业软件算法、数值计算和客户合同交付。\n蒋寻涯：拟兼职加盟，任科学顾问，负责半导体光电仿真、全波电磁建模、L-I-V 测试验证技术路线、产品技术路线评审，并协调对接种子客户和产业资源。其为复旦大学研究员、博士生导师，曾入选中科院百人计划，并推进 EastWave 国产电磁仿真软件产业化。\n团队结构：项目早期采用小而精团队。全职核心为申报人万振文，兼职顾问为蒋寻涯；在完成首批客户验证后，再招聘软件工程师、应用工程师和销售人员。"
)
set_cell_text(
    t3.cell(14, 0),
    "商业模式分三层：1. 联合验证/模型适配服务，按器件型号或项目收费，用于客户数据接入、模型适配和验证报告；2. SaaS 订阅，按账号数、器件型号数和数据量收费，适合研发团队持续使用；3. 私有化部署/API 授权，面向头部光模块、激光芯片和测试装备企业，收取部署费、年度授权费和维护费。市场拓展先以蒋寻涯教授协调的种子客户及苏州光通信产业链为入口，完成数据样例验证，再推进付费试点和年度授权。"
)
set_cell_text(
    t3.cell(16, 0),
    "第一年：完成苏州公司注册，万振文作为全职核心成员落地，蒋寻涯作为兼职科学顾问参与技术路线、测试验证和种子客户协调；完成 TLaser V1.0 工程版，接入至少 1 家客户真实 L-I-V 数据，主营收入 80 万元，新增就业 5 人。\n第二年：形成多器件型号模型库和 SaaS 试点版，完成 3 家以上付费试点，主营收入 350 万元，团队 12 人。\n第三年：形成商业版和私有化部署能力，完成 6 类以上器件型号验证，主营收入 900 万元，团队 25 人。三年累计受理发明专利 5 项、PCT 2 项、软著 6 项，力争高企培育/申报和园区/市级科技项目支持。"
)

replace_para_start(
    doc,
    "TLaser 面向光通讯半导体激光器研发与量产测试环节",
    "TLaser 面向光通讯半导体激光器研发与量产测试环节，建设“数字孪生 + 在线标定 + 批次诊断”的工业软件平台。项目负责人万振文为中科院博士、国家级人才，具备科研建模、工业软件、数值算法和创业交付经验；蒋寻涯教授拟兼职加盟，任科学顾问，负责光电仿真技术路线、L-I-V 测试验证技术路线和种子客户协调对接。\n\n项目背景是光模块和激光芯片企业在高速产品迭代中存在设计仿真耗时、测试数据分散、器件批次漂移难解释、失效分析依赖人工经验等问题。目标产品 TLaser 可接入客户 L-I-V 测试曲线，反演内部物性参数，预测不同设计参数下的输出功率、效率和空间剖面，并形成批次一致性诊断报告。应用领域包括光通讯激光器芯片设计、光模块研发测试、封装筛选、失效分析、量产质量诊断和测试装备算法嵌入。"
)
replace_para_start(
    doc,
    "本项目拟以苏州麦慎数字科技有限公司作为园区产业化主体",
    "本项目拟以苏州麦慎数字科技有限公司作为园区产业化主体，注册时间、注册资本和股权结构以工商登记为准。公司定位为光电子 AI 工业软件企业，首个产品为 TLaser 光通讯激光器数字孪生与在线标定平台。当前项目尚未形成规模化营业收入，但已有可运行原型、自动化验证流程、用户手册和演示材料，并具备历史合同交付和发明专利背书。\n\n落地苏州后，公司第一阶段收入来自联合验证和模型适配服务，第二阶段形成 SaaS 订阅和私有化部署授权，第三阶段扩展到光电子器件数字孪生模型库。早期团队保持小规模：万振文负责全职经营、算法和产品，蒋寻涯教授以兼职科学顾问方式参与技术路线、测试验证和种子客户协调；根据试点进度招聘软件工程、应用工程和销售人员。"
)
replace_para_start(
    doc,
    "项目定位于光电子工业软件和 AI+半导体器件研发工具",
    "项目定位于光电子工业软件和 AI+半导体器件研发工具，切入点是光通讯半导体激光器的设计、测试和量产一致性分析。优先客户包括苏州工业园区及苏州光通信产业链企业：高端光模块企业、激光芯片/泵浦模块企业、光学元件企业、硅光/光电芯片企业和智能检测装备企业。蒋寻涯教授负责协助对接光电仿真、光芯片和测试验证相关种子客户。\n\n市场中已有通用 CAE、EDA、光学仿真和测试设备软件。TLaser 的目标场景为“光通讯激光器实测数据闭环数字孪生”，通过 L-I-V 数据接入、模型校准和批次诊断报告服务研发与测试部门。"
)
replace_para_start(
    doc,
    "商业模式清晰分为三步",
    "商业模式分三步。第一步，联合验证和模型适配服务：选择种子客户，用客户历史 L-I-V 数据训练/校准模型并输出诊断报告，按项目收费。第二步，SaaS 订阅：形成标准云端版本，按账号、器件型号、数据量和模型库收费。第三步，私有化部署/API 授权：面向头部光模块、激光芯片和测试装备企业，收取部署费、年度授权费和维护费。\n\n市场拓展路径为“样例诊断-付费试点-年度授权”。先用客户已有测试数据形成样例报告，再通过模型适配绑定客户器件型号，最后将模型库、API 和私有部署纳入客户研发/测试流程。第一阶段以万振文直接对接和蒋寻涯协调种子客户为主，第二年建立应用工程和销售团队。"
)
replace_para_start(
    doc,
    "3-5 年阶段目标围绕产品、客户、团队和知识产权展开",
    "3-5 年阶段目标围绕产品、客户、团队和知识产权展开。第一年完成苏州公司注册，形成万振文全职负责、蒋寻涯兼职科学顾问支持的小团队结构，完成 TLaser V1.0 工程版、真实 L-I-V 数据接入、1 家以上客户联合验证，主营收入约 80 万元，申请 1 项发明专利、2 项软著，新增就业 5 人。第二年完成多型号激光器模型库、SaaS 试点版、3 家以上付费客户，主营收入约 350 万元，申请 2 项发明专利、1 项 PCT、2 项软著，新增就业至 12 人。第三年完成商业版、私有化部署和 API 接口，服务不少于 6 类器件型号，主营收入约 900 万元，新增就业至 25 人，累计形成 5 项发明专利受理、2 项 PCT、6 项软著。\n\n第四至第五年从激光器扩展到硅光调制器、探测器、光模块热电可靠性和封装测试数字孪生，形成光电子器件模型库和行业案例，争取高新技术企业、专精特新培育和新一轮股权融资。"
)

for table in doc.tables:
    if len(table.rows) >= 54 and "顾俊" in table_text(table):
        for idx in range(53, 26, -1):
            delete_row(table, idx)
    if len(table.rows) >= 27 and "蒋寻涯" in table_text(table):
        set_cell_text(table.cell(8, 3), "□全职  ☑兼职（科学顾问）")
        set_cell_text(table.cell(8, 8), "20%-30%")
        set_cell_text(table.cell(9, 3), "复旦大学在职，拟以科学顾问方式兼职加盟，负责光电仿真技术路线评审、L-I-V 测试验证路线设计、种子客户协调对接和产业资源协同。")
        set_cell_text(
            table.cell(26, 0),
            "就业经历和业绩：蒋寻涯教授为复旦大学研究员、博士生导师。公开资料显示，其曾在美国 Ames 国家实验室、Dicon Fiberoptics、MIT 电子研究所、中科院上海微系统所、复旦大学等机构从事光子学、光电器件、电磁与光电仿真研究；研究方向包括激光/LED、光子晶体、超材料、特异介质、非线性和无序体系等。公开资料显示，其发表 SCI 论文逾百篇，被引用 3000 余次，出版 Springer 学术专著，拥有多项发明专利，并曾入选中科院百人计划、上海杰出引进人才、浦江计划、浙江省千人计划等人才项目。\n\n与本项目关联：蒋寻涯教授创办上海东峻信息科技有限公司，推进国产电磁/光电仿真软件 EastWave 产业化。其在本项目中拟兼职加盟，任科学顾问，负责光电仿真技术路线评审、激光器 L-I-V 测试验证路线、测试数据与物理模型一致性评审，并协调对接种子客户和产业资源。"
        )

doc.save(str(OUT))
print(str(OUT))
print("summary_chars", len(summary))
