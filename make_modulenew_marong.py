from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.table import Table


SRC = Path("ModuleNew_0813.docx")
OUT = Path("ModuleNew_0813_MaRong.docx")


def set_cell(cell, text):
    cell.text = text


def rc(table, row, col):
    return table.rows[row].cells[col]


def replace_in_paragraphs(doc, replacements):
    for paragraph in doc.paragraphs:
        text = paragraph.text
        for old, new in replacements:
            if old in text:
                text = text.replace(old, new)
        if text != paragraph.text:
            paragraph.text = text


def replace_in_cells(doc, replacements):
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text
                for old, new in replacements:
                    if old in text:
                        text = text.replace(old, new)
                if text != cell.text:
                    cell.text = text


def fill_ma_rong(table):
    set_cell(rc(table, 0, 1), "马戎")
    set_cell(rc(table, 0, 5), "提交前补录")
    set_cell(rc(table, 1, 1), "男")
    set_cell(rc(table, 1, 5), "中国")
    set_cell(rc(table, 2, 5), "深圳企业经营管理与商务拓展")
    set_cell(rc(table, 3, 1), "本科")
    set_cell(rc(table, 3, 5), "提交前补录")
    set_cell(rc(table, 4, 1), "工学/管理")
    set_cell(rc(table, 4, 5), "自动化装备、电子开发、客户商务")
    set_cell(rc(table, 5, 3), "□身份证  □护照号（提交前补录）")
    set_cell(rc(table, 5, 8), "提交前补录")
    set_cell(rc(table, 6, 3), "□有  ☑无（公开资料未见海外经历）")
    set_cell(rc(table, 6, 8), "深圳")
    set_cell(rc(table, 7, 3), "□是  ☑否（拟按项目需要来苏州/广州/深圳）")
    set_cell(rc(table, 7, 8), "按项目进度安排")
    set_cell(rc(table, 8, 3), "☑全职  □兼职")
    set_cell(rc(table, 8, 8), "")
    set_cell(
        rc(table, 9, 3),
        "拟作为公司全职商务负责人，负责客户开发、试点合作、销售渠道、合同谈判、回款管理和大湾区/华南产业资源导入。",
    )
    set_cell(rc(table, 10, 3), "提交前补录")
    set_cell(rc(table, 10, 8), "提交前补录")
    set_cell(rc(table, 13, 0), "1997.12")
    set_cell(rc(table, 13, 2), "至今")
    set_cell(rc(table, 13, 4), "深圳市光彩凯宜电子开发有限公司")
    set_cell(rc(table, 13, 7), "法定代表人、执行董事&总经理")
    for r in [14, 15, 16]:
        for c in [0, 2, 4, 7]:
            set_cell(rc(table, r, c), "")
    set_cell(rc(table, 20, 0), "提交前补录")
    set_cell(rc(table, 20, 2), "提交前补录")
    set_cell(rc(table, 20, 4), "提交前补录")
    set_cell(rc(table, 20, 7), "提交前补录")
    set_cell(rc(table, 20, 9), "提交前补录")
    set_cell(
        table.rows[26].cells[0],
        "就业经历和业绩：公开资料显示，马戎先生为深圳市光彩凯宜电子开发有限公司法定代表人、执行董事&总经理。该公司成立于1997年，是自动化装备、老化测试装备、精密工装治具、电子电路开发设计制造服务商，公开项目资料列示其主要客户包括华为、H3C、浪潮信息、曙光信息等ICT领域企业。马戎先生长期在深圳经营电子开发和自动化测试装备企业，具备面向大型ICT客户的商务开拓、客户服务、项目交付和供应链协同经验。\n\n与本项目关联：TLaser早期商业化需要快速获得光模块、激光芯片、测试装备和ICT客户试点。马戎先生拟担任商务负责人，负责苏州、深圳/华南及大湾区客户资源对接、试点项目商务谈判、销售渠道建设、合同回款和测试装备合作伙伴导入，补足团队在市场开拓和商业闭环方面的能力。",
    )


doc = Document(SRC)

summary_team = (
    "核心团队由计算物理、工业软件、半导体光电仿真和商务拓展人员组成。\n"
    "CEO，万振文，中科院博士、美国 MBL 博士后、厦门大学副教授、丹麦 DMI 资深研究员、国家级人才“启明计划”入选者，麦慎数字科技创始人。其长期从事多物理场数值模式、CAE 算法和工业软件开发，完成 EML 高频电磁仿真、MySim 多物理场仿真等软件原型，并服务清华大学深圳研究生院、天津科技大学、上海市气象局、北京海基、上海应尧、苏州 ShonCloud、EastWave、Ansys 中国等高校、科研机构和企业客户。已授权发明专利 3 项，具备将复杂物理模型工程化为软件产品的经验。\n"
    "商务负责人，马戎，拟全职加盟。公开资料显示，其为深圳市光彩凯宜电子开发有限公司法定代表人、执行董事&总经理，长期经营自动化装备、老化测试装备、精密工装治具和电子电路开发设计制造业务，服务华为、H3C、浪潮信息、曙光信息等 ICT 客户，具备深圳/华南大客户商务、项目交付和供应链协同经验。\n"
    "科学顾问，蒋寻涯，复旦大学教授、博士生导师，西湖大学光电研究院特聘教授，中科院百人计划入选者，EastWave 电磁仿真软件第一产权人，杭州曦岭科技创始人。其负责半导体光电仿真技术路线、L-I-V 测试验证方法和种子客户资源协同。\n"
    "杭州曦岭科技侧重半导体物理仿真，苏州麦慎数字科技（拟）侧重光通讯激光器7D即时预测、L-I-V数据标定和SaaS工程化；马戎补强深圳/大湾区客户开拓和测试装备产业资源导入。"
)
doc.tables[3].rows[3].cells[0].text = summary_team

replacements = [
    (
        "项目负责人万振文为中科院博士、国家级人才，具备科研建模、工业软件、数值算法和创业交付经验；蒋寻涯教授兼职加盟，任科学顾问，负责光电仿真技术路线、L-I-V测试验证技术路线和种子客户协调对接。",
        "项目负责人万振文为中科院博士、国家级人才，具备科研建模、工业软件、数值算法和创业交付经验；马戎先生拟全职担任商务负责人，负责客户试点、渠道建设、合同谈判和深圳/大湾区产业资源导入；蒋寻涯教授兼职加盟，任科学顾问，负责光电仿真技术路线、L-I-V测试验证技术路线和种子客户协调对接。",
    ),
    (
        "早期团队保持小规模：万振文负责全职经营、算法和产品，蒋寻涯教授以兼职科学顾问方式参与技术路线、L-I-V测试验证和种子客户协调；根据试点进度招聘软件工程、应用工程和销售人员。",
        "早期团队保持小规模：万振文负责全职经营、算法和产品，马戎拟全职负责商务拓展、客户试点、销售渠道和合同回款，蒋寻涯教授以兼职科学顾问方式参与技术路线、L-I-V测试验证和种子客户协调；根据试点进度招聘软件工程、应用工程和销售人员。",
    ),
    (
        "蒋寻涯教授负责协助对接光电仿真、光芯片和测试验证相关种子客户。",
        "蒋寻涯教授负责协助对接光电仿真、光芯片和测试验证相关种子客户；马戎负责对接深圳/华南ICT客户、光模块测试装备伙伴和大湾区商务渠道。",
    ),
    (
        "第一阶段以万振文直接对接和蒋寻涯协调种子客户为主，第二年建立应用工程和销售团队。",
        "第一阶段以万振文直接对接、马戎推进商务试点、蒋寻涯协调种子客户为主，第二年建立应用工程和销售团队。",
    ),
    (
        "第一年完成苏州公司注册，形成万振文全职负责、蒋寻涯兼职科学顾问支持的小团队结构，",
        "第一年完成苏州公司注册，形成万振文全职负责、马戎全职商务拓展、蒋寻涯兼职科学顾问支持的小团队结构，",
    ),
    (
        "第一年：完成苏州公司注册，万振文作为全职核心成员落地，蒋寻涯作为兼职科学顾问参与技术路线、测试验证和种子客户协调；",
        "第一年：完成苏州公司注册，万振文作为全职技术/产品核心成员落地，马戎作为全职商务负责人推进客户试点和渠道建设，蒋寻涯作为兼职科学顾问参与技术路线、测试验证和种子客户协调；",
    ),
    (
        "新增商务/应用工程人员1人。",
        "新增商务/应用工程人员1人。马戎负责深圳/华南及大湾区客户资源导入。",
    ),
]
replace_in_paragraphs(doc, replacements)
replace_in_cells(doc, replacements)

# Duplicate the current team-member table and fill the copy with Ma Rong.
jiang_table = doc.tables[15]
new_tbl = deepcopy(jiang_table._tbl)
jiang_table._tbl.addnext(new_tbl)
ma_table = Table(new_tbl, jiang_table._parent)
fill_ma_rong(ma_table)

doc.save(OUT)
print(f"saved={OUT}")
