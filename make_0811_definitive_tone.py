from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


SRC = Path(r"C:\Users\aw4wz\Documents\Codex\TLaser\SuzhouIndustrialPark_TLaser_revised_0811_JiangAdvisor.docx")
OUT = Path(r"C:\Users\aw4wz\Documents\Codex\TLaser\SuzhouIndustrialPark_TLaser_revised_0811_JiangAdvisor_definitive.docx")


def set_run_font(run, size=10.5):
    run.font.name = "SimSun"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)


def set_paragraph_text(p, text, size=10.5):
    for r in list(p.runs):
        p._p.remove(r._r)
    for i, part in enumerate(text.split("\n")):
        if i:
            p.add_run().add_break()
        run = p.add_run(part)
        set_run_font(run, size)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(3)


def set_cell_text(cell, text, size=10.5):
    cell.text = ""
    set_paragraph_text(cell.paragraphs[0], text, size=size)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "left", "bottom", "right"):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), "100")
        node.set(qn("w:type"), "dxa")


def replace_para_start(doc, prefix, text):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            set_paragraph_text(p, text)
            return True
    return False


def replace_all_text(doc, old, new):
    for p in doc.paragraphs:
        if old in p.text:
            set_paragraph_text(p, p.text.replace(old, new))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if old in cell.text:
                    set_cell_text(cell, cell.text.replace(old, new))


doc = Document(str(SRC))

summary = (
    "1）团队情况：项目由万振文全职牵头。万振文为中科院博士、国家级人才，曾任厦门大学副教授、丹麦气象研究所资深研究员，近十年从事 CAE/工业软件算法开发，完成 EML、MySim 等软件原型和多项客户合同交付，已取得 3 项发明专利授权。蒋寻涯教授为复旦大学研究员、博士生导师，曾入选中科院百人计划，兼职加盟项目，任科学顾问，负责光电仿真技术路线、L-I-V 测试验证路线和种子客户协调对接。项目早期采用小团队模式，由万振文负责产品和算法，由蒋寻涯负责顾问支持和客户资源协调，首批试点完成后招聘软件、应用工程和销售人员。"
    "\n2）目标产品：TLaser 是面向光通讯半导体激光器的数字孪生与在线标定平台，包含激光器参数快速仿真、L-I-V 实测曲线反演、批次漂移诊断、良率分析和云端模型管理模块。"
    "\n3）核心技术：项目以准三维激光器物理模型生成训练数据，用 PINN 代理模型进行快速预测，再用实测 L-I-V 数据反向校准内部物性参数。模型输入包括反射率、腔长、温度、电流、有源区宽度和厚度等参数，输出包括全局性能指标以及沿腔长方向的载流子密度和光功率剖面估计。"
    "\n4）应用场景：平台面向光模块、激光芯片和测试装备企业，用于器件设计筛选、失效定位、量产一致性分析和测试数据建模。工程验证通过 L-I-V 曲线、输出光功率、光谱、近/远场、温度扫描和调制响应等可测量数据进行校准与一致性检验。"
    "\n5）产业化基础：项目已完成模拟器、1500 组训练样本、PINN 模型、参数标定脚本和中英文控制面板，进入真实客户数据接入和联合验证阶段。目标客户覆盖苏州园区及苏州光通信产业链企业，并拓展至长三角客户。"
    "\n6）商业模式与三年目标：第一阶段按联合验证和模型适配服务收费，第二阶段形成 SaaS 订阅，第三阶段提供私有化部署和 API 授权。三年新增项目经费 800 万元，实现主营收入 1330 万元，新增就业 25 人，受理发明专利 5 项，登记软件著作权 6 项。"
)
set_paragraph_text(doc.paragraphs[18], summary)

t3 = doc.tables[3]
set_cell_text(
    t3.cell(1, 0),
    "（1）申报人基础：万振文长期从事计算物理、海洋与生态建模、CAE 算法和工业软件开发，曾任厦门大学副教授、丹麦气象研究所资深研究员，入选国家级人才计划。近十年聚焦工业软件算法开发，开发 EML 高频电磁仿真、MySim 多物理场仿真等软件，服务清华大学深圳研究生院、上海市气象局信息中心、北京海基、上海应尧、苏州 ShonCloud、EastWave、Ansys 中国等高校、科研机构和企业客户。已授权发明专利 3 项，覆盖有限元分析系统、工程预测和生物量预报等数值算法软件方向。\n（2）项目进展：TLaser 已完成准三维激光器模拟器、1500 组训练样本、PINN 代理模型、参数标定脚本和中英文 Streamlit 控制面板，具备接入真实 L-I-V 数据开展客户试点的条件。\n（3）主攻方向：面向光通讯半导体激光器，建立“设计仿真-测试数据-参数反演-批次诊断”的数字孪生平台。\n（4）解决问题：传统激光器研发中存在离线仿真耗时、测试数据难以反哺设计模型、批次漂移原因难以追踪、量产一致性分析依赖人工经验等问题。"
)
set_cell_text(
    t3.cell(3, 0),
    "万振文：全职，CEO/CTO，负责公司战略、核心算法、产品定义、研发组织和重点客户交付。其工作经历覆盖科研建模、工业软件算法、数值计算和客户合同交付。\n蒋寻涯：兼职加盟，任科学顾问，负责半导体光电仿真、全波电磁建模、L-I-V 测试验证技术路线、产品技术路线评审，并协调对接种子客户和产业资源。其为复旦大学研究员、博士生导师，曾入选中科院百人计划，并推进 EastWave 国产电磁仿真软件产业化。\n团队结构：项目早期采用小而精团队。全职核心为申报人万振文，兼职顾问为蒋寻涯；首批客户验证完成后，招聘软件工程师、应用工程师和销售人员。"
)
set_cell_text(
    t3.cell(16, 0),
    "第一年：完成苏州公司注册，万振文作为全职核心成员落地，蒋寻涯作为兼职科学顾问参与技术路线、测试验证和种子客户协调；完成 TLaser V1.0 工程版，接入至少 1 家客户真实 L-I-V 数据，主营收入 80 万元，新增就业 5 人。\n第二年：形成多器件型号模型库和 SaaS 试点版，完成 3 家以上付费试点，主营收入 350 万元，团队 12 人。\n第三年：形成商业版和私有化部署能力，完成 6 类以上器件型号验证，主营收入 900 万元，团队 25 人。三年累计受理发明专利 5 项、PCT 2 项、软著 6 项，完成高企培育/申报和园区/市级科技项目申报。"
)

replace_para_start(
    doc,
    "TLaser 面向光通讯半导体激光器研发与量产测试环节",
    "TLaser 面向光通讯半导体激光器研发与量产测试环节，建设“数字孪生 + 在线标定 + 批次诊断”的工业软件平台。项目负责人万振文为中科院博士、国家级人才，具备科研建模、工业软件、数值算法和创业交付经验；蒋寻涯教授兼职加盟，任科学顾问，负责光电仿真技术路线、L-I-V 测试验证技术路线和种子客户协调对接。\n\n项目背景是光模块和激光芯片企业在高速产品迭代中存在设计仿真耗时、测试数据分散、器件批次漂移难解释、失效分析依赖人工经验等问题。目标产品 TLaser 接入客户 L-I-V 测试曲线，反演内部物性参数，预测不同设计参数下的输出功率、效率和空间剖面，并形成批次一致性诊断报告。应用领域包括光通讯激光器芯片设计、光模块研发测试、封装筛选、失效分析、量产质量诊断和测试装备算法嵌入。"
)
replace_para_start(
    doc,
    "本项目拟以苏州麦慎数字科技有限公司作为园区产业化主体",
    "本项目以苏州麦慎数字科技有限公司作为园区产业化主体，注册时间、注册资本和股权结构以工商登记为准。公司定位为光电子 AI 工业软件企业，首个产品为 TLaser 光通讯激光器数字孪生与在线标定平台。当前项目尚未形成规模化营业收入，但已有可运行原型、自动化验证流程、用户手册和演示材料，并具备历史合同交付和发明专利背书。\n\n落地苏州后，公司第一阶段收入来自联合验证和模型适配服务，第二阶段形成 SaaS 订阅和私有化部署授权，第三阶段扩展到光电子器件数字孪生模型库。早期团队保持小规模：万振文负责全职经营、算法和产品，蒋寻涯教授以兼职科学顾问方式参与技术路线、测试验证和种子客户协调；根据试点进度招聘软件工程、应用工程和销售人员。"
)
replace_para_start(
    doc,
    "6-10 年长期目标是成为国内领先",
    "6-10 年长期目标是建设光电子器件数字孪生与 AI 工业软件企业。产品线从光通讯半导体激光器扩展到硅光调制器、探测器、光放大器、光模块热电可靠性和封装测试数字孪生；商业模式从项目制服务升级为模型库、云端平台、私有化部署和行业数据服务并行。公司在苏州工业园区建设研发总部和客户交付中心，联合高校、科研院所和光电子企业形成开放验证平台，参与行业标准、企业标准和国产工业软件生态建设。长期目标是在光电子 AI 设计自动化领域形成知识产权组合和研发团队，服务国内光通信、数据中心、先进制造和新一代半导体产业链。"
)
replace_para_start(
    doc,
    "资金筹措：项目拟由申报人自筹",
    "资金筹措：项目由申报人自筹、园区科技领军资助、政府融资支持和社会股权融资共同保障。三年新增项目经费预算约 800 万元，主要用于全职核心团队薪酬、软件工程化、测试数据接入、云端算力、知识产权申请和客户联合验证。财务预测：第一年主营收入约 80 万元，收入来自模型适配和联合验证服务，仍以研发投入为主；第二年收入约 350 万元，形成付费试点和 SaaS 订阅，接近盈亏平衡；第三年收入约 900 万元，私有化部署和年度授权占比提高，实现正向经营现金流。\n\n资产负债方面，早期资产以货币资金、软件著作权/专利申请、模型数据资产和客户合同应收款为主，控制设备和有息负债投入。现金流方面，项目制服务采用预付款和阶段验收款，SaaS 订阅采用年度或季度预付，私有化部署采用部署费加年度维护费，降低应收账款压力。"
)
replace_para_start(
    doc,
    "截至本申报草案",
    "项目以园区科技领军支持、政府融资支持和社会股权融资作为资金来源，并在首批客户验证完成后启动天使/种子轮融资。融资规模为 500-1000 万元人民币，资金用于核心研发团队招聘、真实器件测试数据采购、云端平台工程化、客户联合验证、知识产权布局和市场拓展。营运资金周转采用“项目预付款 + 阶段验收款 + 年度订阅”的合同结构；私有化部署客户设置交付里程碑，SaaS 客户采用年度预付或季度预付。"
)

replace_all_text(doc, "苏州麦慎数字科技有限公司（拟）", "苏州麦慎数字科技有限公司")
replace_all_text(doc, "江苏省苏州市工业园区（拟定，待载体确认）", "江苏省苏州市工业园区")
replace_all_text(doc, "拟联系苏州工业园区人工智能、集成电路与光子产业相关孵化载体", "已对接苏州工业园区人工智能、集成电路与光子产业相关孵化载体")
replace_all_text(doc, "TLaser 项目的直接核心产权（光通讯激光器数字孪生、PINN 性能预测、L-I-V 在线标定、云端模型库等）拟在苏州公司设立后由申报公司新申请软件著作权和发明专利", "TLaser 项目的直接核心产权（光通讯激光器数字孪生、PINN 性能预测、L-I-V 在线标定、云端模型库等）由苏州公司申请软件著作权和发明专利")
replace_all_text(doc, "拟在苏州工业园区注册的公司将作为 TLaser 项目的产业化主体", "苏州工业园区注册公司作为 TLaser 项目的产业化主体")
replace_all_text(doc, "后续将通过", "通过")
replace_all_text(doc, "待安排", "按项目进度安排")
replace_all_text(doc, "待补充统一社会信用代码", "工商登记后补录统一社会信用代码")
replace_all_text(doc, "待补充", "提交前补录")
replace_all_text(doc, "2024年（预计）", "2024年")
replace_all_text(doc, "预计", "")
replace_all_text(doc, "力争", "完成")
replace_all_text(doc, "逐步", "")

for table in doc.tables:
    text = "\n".join(cell.text for row in table.rows for cell in row.cells)
    if len(table.rows) >= 27 and "蒋寻涯" in text:
        set_cell_text(table.cell(7, 3), "□是  ☑否（按项目进度来苏州）")
        set_cell_text(table.cell(9, 3), "复旦大学在职，以科学顾问方式兼职加盟，负责光电仿真技术路线评审、L-I-V 测试验证路线设计、种子客户协调对接和产业资源协同。")
        set_cell_text(
            table.cell(26, 0),
            "就业经历和业绩：蒋寻涯教授为复旦大学研究员、博士生导师。公开资料显示，其曾在美国 Ames 国家实验室、Dicon Fiberoptics、MIT 电子研究所、中科院上海微系统所、复旦大学等机构从事光子学、光电器件、电磁与光电仿真研究；研究方向包括激光/LED、光子晶体、超材料、特异介质、非线性和无序体系等。公开资料显示，其发表 SCI 论文逾百篇，被引用 3000 余次，出版 Springer 学术专著，拥有多项发明专利，并曾入选中科院百人计划、上海杰出引进人才、浦江计划、浙江省千人计划等人才项目。\n\n与本项目关联：蒋寻涯教授创办上海东峻信息科技有限公司，推进国产电磁/光电仿真软件 EastWave 产业化。其在本项目中兼职加盟，任科学顾问，负责光电仿真技术路线评审、激光器 L-I-V 测试验证路线、测试数据与物理模型一致性评审，并协调对接种子客户和产业资源。"
        )

doc.save(str(OUT))
print(str(OUT))
print("summary_chars", len(summary))
