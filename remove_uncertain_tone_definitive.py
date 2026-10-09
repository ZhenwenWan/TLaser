from pathlib import Path

import docx
from docx.oxml.ns import qn


PATH = Path(r"C:\Users\aw4wz\Documents\Codex\TLaser\SuzhouIndustrialPark_TLaser_revised_0811_JiangAdvisor_definitive.docx")


def set_run_font(run, size=None, bold=None):
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.bold = bold
    run.font.name = "仿宋"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "仿宋")


def replace_text(paragraph, replacements):
    text = paragraph.text
    new_text = text
    for old, new in replacements:
        new_text = new_text.replace(old, new)
    if new_text == text:
        return False
    paragraph.clear()
    run = paragraph.add_run(new_text)
    set_run_font(run)
    return True


REPLACEMENTS = [
    ("拟创办", "创办"),
    ("(拟)创办", "创办"),
    ("拟注册", "注册"),
    ("拟认缴", "认缴"),
    ("拟实缴", "实缴"),
    ("拟持股", "持股"),
    ("拟解决", "解决"),
    ("拟定，", ""),
    ("拟定", ""),
    ("拟注册主体", "注册主体"),
    ("本项目拟注册公司", "本项目注册公司"),
    ("待工商信息确认", "工商信息确认后承接"),
    ("待", ""),
    ("模拟器", "仿真器"),
]


doc = docx.Document(PATH)
for paragraph in doc.paragraphs:
    replace_text(paragraph, REPLACEMENTS)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                replace_text(paragraph, REPLACEMENTS)

doc.save(PATH)
print(PATH)
