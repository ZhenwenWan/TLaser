from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.shared import Pt


PATH = Path(r"C:\Users\aw4wz\Documents\Codex\TLaser\SuzhouIndustrialPark_TLaser_revised_0813_ChinaPhotonics.docx")


def set_font(run, size=10.5):
    run.font.name = "仿宋"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "仿宋")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)


def replace_cell(cell, text):
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    p = cell.paragraphs[0]
    p.clear()
    run = p.add_run(text)
    set_font(run)


doc = docx.Document(PATH)
team_advantage = (
    "核心团队采用“小而精”结构，由计算物理、工业软件和半导体光电仿真专家组成。\n"
    "CEO万振文，中科院博士、美国MBL博士后、厦门大学副教授、丹麦DMI资深研究员、国家级人才“启明计划”入选者，麦慎数字科技创始人。其长期从事多物理场数值模式、CAE算法和工业软件开发，完成EML高频电磁仿真、MySim多物理场仿真等软件原型，并服务清华大学深圳研究生院、天津科技大学、上海市气象局、北京海基、上海应尧、苏州ShonCloud、EastWave、Ansys中国等高校、科研机构和企业客户。已授权发明专利3项，具备将复杂物理模型工程化为软件产品的经验。\n"
    "科学顾问蒋寻涯，复旦大学教授、博士生导师，西湖大学光电研究院特聘教授，中科院百人计划入选者，EastWave电磁仿真软件第一产权人，杭州曦岭科技创始人。其负责半导体光电仿真技术路线、L-I-V测试验证方法和种子客户资源协同。\n"
    "团队分工与TLaser产品定位一致：万振文全职负责经营、算法、产品和交付；蒋寻涯兼职加盟，负责科学顾问、测试验证路线和种子客户协调。苏州麦慎数字科技侧重光通讯激光器7D即时预测、L-I-V数据标定和SaaS工程化，产品边界清晰。"
)

replace_cell(doc.tables[3].rows[1].cells[0], team_advantage)
doc.save(PATH)

text = "\n".join(
    [p.text for p in doc.paragraphs]
    + [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
)
print(PATH)
print({k: text.count(k) for k in ["拟", "待", "预计", "力争", "顾俊", "马戎", "秦三团", "中国光芯", "1624亿元"]})
