from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt


SRC = Path(r"C:\Users\aw4wz\Documents\Codex\TLaser\SuzhouIndustrialPark_TLaser_revised_0810_GuJun.docx")
OUT = Path(r"C:\Users\aw4wz\Documents\Codex\TLaser\SuzhouIndustrialPark_TLaser_revised_0810_GuJun_summary_refined.docx")


def set_run_font(run, size=10.5):
    run.font.name = "SimSun"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)


def set_paragraph_text(p, text):
    for r in list(p.runs):
        p._p.remove(r._r)
    for i, part in enumerate(text.split("\n")):
        if i:
            p.add_run().add_break()
        run = p.add_run(part)
        set_run_font(run)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(3)


summary = (
    "1）团队情况：项目由万振文全职牵头。万振文为中科院博士、国家级人才，曾任厦门大学副教授、丹麦气象研究所资深研究员，近十年从事 CAE/工业软件算法开发，完成 EML、MySim 等软件原型和多项客户合同交付，已取得 3 项发明专利授权。顾俊拟全职担任商务负责人，公开资料显示其曾任苏州舜云工程软件有限公司仿真技术经理，负责动力系统模块解决方案研发及推广，并完成汽车行业多个头部企业热管理分析技术移交。蒋寻涯教授为复旦大学研究员、博士生导师，曾入选中科院百人计划，拟作为兼职首席科学顾问负责光电仿真、L-I-V 测试验证技术路线和产业资源协同。"
    "\n2）目标产品：TLaser 是面向光通讯半导体激光器的数字孪生与在线标定平台，包含激光器参数快速仿真、L-I-V 实测曲线反演、批次漂移诊断、良率分析和云端模型管理模块。"
    "\n3）核心技术：项目以准三维激光器物理模型生成训练数据，用 PINN 代理模型进行快速预测，再用实测 L-I-V 数据反向校准内部物性参数。当前模型输入包括反射率、腔长、温度、电流、有源区宽度和厚度等参数，输出包括全局性能指标以及沿腔长方向的载流子密度和光功率剖面估计。"
    "\n4）应用场景：平台面向光模块、激光芯片和测试装备企业，用于器件设计筛选、失效定位、量产一致性分析和测试数据建模。工程验证主要通过 L-I-V 曲线、输出光功率、光谱、近/远场、温度扫描和调制响应等可测量数据进行校准与一致性检验。"
    "\n5）产业化基础：项目已完成模拟器、1500 组训练样本、PINN 模型、参数标定脚本和中英文控制面板，可进入真实客户数据接入和联合验证阶段。目标客户优先覆盖苏州园区及苏州光通信产业链企业，并逐步拓展至长三角和深圳/华南客户。"
    "\n6）商业模式与三年目标：第一阶段按联合验证和模型适配服务收费，第二阶段形成 SaaS 订阅，第三阶段提供私有化部署和 API 授权。三年计划新增项目经费 800 万元，实现主营收入 1330 万元，新增就业 25 人，受理发明专利 5 项，登记软件著作权 6 项。"
)

doc = Document(str(SRC))
doc.paragraphs[18].text
set_paragraph_text(doc.paragraphs[18], summary)
doc.save(str(OUT))

print("old_doc", SRC)
print("new_doc", OUT)
print("summary_chars", len(summary))
